#!/usr/bin/env python3
"""Portable, deterministic native-PPTX writer for DeliveryIR v1.

The production path is intentionally independent from Codex/Node runtimes.
It requires the pinned ``python-pptx==1.0.2`` distribution (and Pillow, which
is a python-pptx dependency).  QA PNGs are deterministic previews generated
from the exact same geometry ledger; they are audit artifacts and are never
embedded into the editable PPTX.
"""

from __future__ import annotations

import argparse
from datetime import date
import json
import math
import os
from pathlib import Path
import sys
import tempfile
import unicodedata
from typing import Any

sys.dont_write_bytecode = True

EXPECTED_PYTHON_PPTX = "1.0.2"
EXIT_HOLD = 3
WIDTH = 1280
HEIGHT = 720
EMU_PER_PIXEL = 9525
CLOSED_SLIDE_SIZE = {
    "cx": str(WIDTH * EMU_PER_PIXEL),
    "cy": str(HEIGHT * EMU_PER_PIXEL),
    "type": "screen16x9",
}
PAGE = {"left": 64, "top": 48, "width": 1152, "height": 620}
FONT = "Microsoft YaHei"
DISPLAY_SERIES = "诉讼可视化简报"
DISPLAY_EYEBROW = "诉讼可视化简报 · 客户审阅材料"
DISPLAY_REVIEW_STATE = "审阅稿"
DISPLAY_CONFIDENTIALITY = {
    "CONFIDENTIAL": "机密",
    "ATTORNEY_WORK_PRODUCT": "律师工作成果",
    "INTERNAL_REVIEW_ONLY": "仅供内部审阅",
}
NOTES_SOURCE_HEADING = "【来源】"
NOTES_BOUNDARY_HEADING = "【使用边界】"
NOTES_BOUNDARY_TEXT = "审阅稿；未自动取得客户放行。"
SEMANTIC_TOKEN = "SEMITEM"
TEXT_ROLE_TOKEN = "TEXTROLE"
POINT_TO_LAYOUT_PIXEL = 1 / 0.75
TYPE_PT = {"DECK": 50.25, "SLIDE": 35.25, "MID": 24.0, "BODY": 16.5, "AUX": 9.0}
TIMELINE_DATE_LEFT = 64
TIMELINE_DATE_WIDTH = 172
TIMELINE_GUIDE_LEFT = 256
TIMELINE_MARKER_LEFT = 245
TIMELINE_TEXT_LEFT = 286
TIMELINE_TEXT_WIDTH = 904
COLORS = {
    "ink": "152033",
    "muted": "667085",
    "line": "D7DEE8",
    "wash": "F4F7FA",
    "blue": "2F64C6",
    "pale_blue": "EAF1FD",
    "amber": "A15C12",
    "pale_amber": "FFF3DD",
    "white": "FFFFFF",
}
TIMELINE_PROFILES = {"a", "b"}
TIMELINE_AXIS_LEFT = 224
TIMELINE_AXIS_RIGHT = 1056
TIMELINE_AXIS_Y = 382

_SOURCE_LAYOUT_HEADERS = {
    "canonical-v1": {
        "A01": ["主体", "角色", "要点", "来源"],
        "A04": ["项目", "金额（元）", "口径"],
        "A06": ["要件", "事实", "证据", "缺口", "状态"],
        "A08": ["阶段", "动作", "出口"],
        "A09": ["编号", "类别", "内容", "解除条件"],
    },
    "legacy-frozen07-v1": {
        "A01": ["角色", "主体", "说明"],
        "A04": ["口径", "项目", "金额（万元）", "说明"],
        "A06": ["要件", "事实主张", "对应证据", "状态", "备注"],
        "A08": ["阶段", "主要动作", "时点"],
        "A09": ["风险事项", "等级", "说明", "处置"],
    },
}


def visible_table_headers(source_schema: Any) -> dict[str, list[str]]:
    """Validate the closed IR schema without sacrificing single-file portability."""

    if not isinstance(source_schema, dict):
        raise ValueError("DeliveryIR缺少源布局契约")
    variant = source_schema.get("variant")
    headers = _SOURCE_LAYOUT_HEADERS.get(variant)
    expected = {
        "schema": "frozen07-source-schema/1.0",
        "variant": variant,
        "table_headers": headers,
    }
    if headers is None or source_schema != expected:
        raise ValueError("DeliveryIR源布局表头与变体不一致")
    # Render the validated IR values, not an implicit canonical fallback.
    ir_headers = source_schema["table_headers"]
    return {
        "matrix": list(ir_headers["A06"]),
        "amounts": list(ir_headers["A04"]),
        "stages": list(ir_headers["A08"]),
        "risks": list(ir_headers["A09"]),
    }


class DependencyUnavailable(RuntimeError):
    pass


def dependency_report() -> dict[str, Any]:
    if os.environ.get("S2_PPTX_FORCE_UNAVAILABLE") == "1":
        raise DependencyUnavailable("python-pptx dependency forced unavailable")
    try:
        import pptx
        import PIL
    except ImportError as exc:
        raise DependencyUnavailable("python-pptx/Pillow dependency unavailable") from exc
    if pptx.__version__ != EXPECTED_PYTHON_PPTX:
        raise DependencyUnavailable(
            f"python-pptx version mismatch: expected {EXPECTED_PYTHON_PPTX}, got {pptx.__version__}"
        )
    return {
        "status": "pass",
        "engine": "python-pptx",
        "python_pptx_version": pptx.__version__,
        "pillow_version": getattr(PIL, "__version__", "unknown"),
        "python_version": ".".join(str(value) for value in sys.version_info[:3]),
        "codex_runtime_required": False,
        "node_runtime_required": False,
    }


def px(value: float):
    from pptx.util import Emu

    return Emu(int(round(value * EMU_PER_PIXEL)))


def _close_presentation_default_text_style(presentation: Any) -> None:
    """Replace the pinned template's English/theme-font presentation defaults."""

    from pptx.oxml.ns import qn

    style_nodes = presentation._element.findall(qn("p:defaultTextStyle"))
    if len(style_nodes) != 1:
        raise ValueError("presentation defaultTextStyle must be unique")
    style = style_nodes[0]
    if style.attrib or [child.tag for child in style] != [
        qn("a:defPPr"),
        *[qn(f"a:lvl{level}pPr") for level in range(1, 10)],
    ]:
        raise ValueError("presentation defaultTextStyle skeleton drifted")
    default_paragraphs = style.findall(qn("a:defPPr"))
    if len(default_paragraphs) != 1:
        raise ValueError("presentation defaultTextStyle defPPr must be unique")
    default_runs = default_paragraphs[0].findall(qn("a:defRPr"))
    if len(default_runs) != 1:
        raise ValueError("presentation defaultTextStyle defRPr must be unique")
    if (
        default_paragraphs[0].attrib
        or list(default_paragraphs[0]) != default_runs
        or default_runs[0].attrib != {"lang": "en-US"}
        or list(default_runs[0])
    ):
        raise ValueError("presentation defaultTextStyle default run drifted")
    default_runs[0].set("lang", "zh-CN")

    for level in range(1, 10):
        paragraphs = style.findall(qn(f"a:lvl{level}pPr"))
        if len(paragraphs) != 1:
            raise ValueError(f"presentation defaultTextStyle level {level} must be unique")
        expected_paragraph_attributes = {
            "marL": str((level - 1) * 457200),
            "algn": "l",
            "defTabSz": "457200",
            "rtl": "0",
            "eaLnBrk": "1",
            "latinLnBrk": "0",
            "hangingPunct": "1",
        }
        if paragraphs[0].attrib != expected_paragraph_attributes:
            raise ValueError(
                f"presentation defaultTextStyle level {level} attributes drifted"
            )
        runs = paragraphs[0].findall(qn("a:defRPr"))
        if len(runs) != 1 or list(paragraphs[0]) != runs:
            raise ValueError(
                f"presentation defaultTextStyle level {level} defRPr must be unique"
            )
        run = runs[0]
        if run.attrib != {"sz": "1800", "kern": "1200"}:
            raise ValueError(
                f"presentation defaultTextStyle level {level} run attributes drifted"
            )
        if [child.tag for child in run] != [
            qn("a:solidFill"),
            qn("a:latin"),
            qn("a:ea"),
            qn("a:cs"),
        ]:
            raise ValueError(
                f"presentation defaultTextStyle level {level} run skeleton drifted"
            )
        solid_fill = run[0]
        if (
            solid_fill.attrib
            or len(solid_fill) != 1
            or solid_fill[0].tag != qn("a:schemeClr")
            or solid_fill[0].attrib != {"val": "tx1"}
            or list(solid_fill[0])
        ):
            raise ValueError(
                f"presentation defaultTextStyle level {level} fill drifted"
            )
        run.set("lang", "zh-CN")
        fonts = []
        for tag, legacy_typeface in (
            ("a:latin", "+mn-lt"),
            ("a:ea", "+mn-ea"),
            ("a:cs", "+mn-cs"),
        ):
            nodes = run.findall(qn(tag))
            if len(nodes) != 1:
                raise ValueError(
                    f"presentation defaultTextStyle level {level} {tag} must be unique"
                )
            if nodes[0].attrib != {"typeface": legacy_typeface} or list(nodes[0]):
                raise ValueError(
                    f"presentation defaultTextStyle level {level} {tag} drifted"
                )
            nodes[0].attrib.clear()
            nodes[0].set("typeface", FONT)
            fonts.append(nodes[0])


def rgb(value: str):
    from pptx.dml.color import RGBColor

    return RGBColor.from_string(value.lstrip("#").upper())


def semantic_name(base: str, ids: list[str] | tuple[str, ...] = ()) -> str:
    return base + "".join(f"|{SEMANTIC_TOKEN}|{value}|" for value in ids)


def text_role_name(base: str, role: str) -> str:
    if role not in TYPE_PT:
        raise ValueError(f"unknown text role {role}")
    return f"{base}|{TEXT_ROLE_TOKEN}|{role}|"


def _font_path() -> str | None:
    candidates = (
        "/System/Library/Fonts/Hiragino Sans GB.ttc",
        "/System/Library/Fonts/STHeiti Medium.ttc",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/opentype/noto/NotoSansCJKsc-Regular.otf",
        "C:/Windows/Fonts/msyh.ttc",
    )
    return next((value for value in candidates if Path(value).is_file()), None)


def display_confidentiality(value: Any) -> str:
    if not isinstance(value, str) or value not in DISPLAY_CONFIDENTIALITY:
        raise ValueError("confidentiality缺少简体中文显示映射")
    return DISPLAY_CONFIDENTIALITY[value]


def _text_weight(value: str) -> float:
    total = 0.0
    for character in value:
        if character.isspace():
            total += 0.45
        elif unicodedata.east_asian_width(character) in {"W", "F", "A"}:
            total += 1.0
        elif character.isupper():
            total += 0.68
        else:
            total += 0.56
    return total


def estimate_lines(text: str, width: float, font_pt: float) -> int:
    usable = max(1.0, width - 18.0)
    glyph_pixels = max(1.0, font_pt * POINT_TO_LAYOUT_PIXEL)
    capacity = max(1.0, usable / glyph_pixels)
    return max(
        1,
        sum(max(1, math.ceil(_text_weight(line) / capacity)) for line in str(text).split("\n")),
    )


def _wrap_for_preview(draw, text: str, font, maximum_width: int) -> str:
    result: list[str] = []
    for source_line in str(text).split("\n"):
        if not source_line:
            result.append("")
            continue
        current = ""
        for character in source_line:
            candidate = current + character
            bounds = draw.textbbox((0, 0), candidate, font=font)
            if current and bounds[2] - bounds[0] > maximum_width:
                result.append(current)
                current = character
            else:
                current = candidate
        result.append(current)
    return "\n".join(result)


class SlideContext:
    def __init__(self, slide, slide_id: str):
        from PIL import Image, ImageDraw

        self.slide = slide
        self.slide_id = slide_id
        self.elements: list[dict[str, Any]] = []
        self.names: set[str] = set()
        self.image = Image.new("RGB", (WIDTH, HEIGHT), "#FFFFFF")
        self.draw = ImageDraw.Draw(self.image)
        self._font_cache: dict[int, Any] = {}

    def register_name(self, name: str) -> None:
        if not name or name in self.names:
            raise ValueError(f"duplicate or empty shape name:{self.slide_id}:{name}")
        self.names.add(name)

    def set_background(self, color: str) -> None:
        fill = self.slide.background.fill
        fill.solid()
        fill.fore_color.rgb = rgb(color)
        self.draw.rectangle((0, 0, WIDTH, HEIGHT), fill=f"#{color}")

    def font(self, points: float):
        from PIL import ImageFont

        pixels = max(8, int(round(points * POINT_TO_LAYOUT_PIXEL)))
        if pixels not in self._font_cache:
            path = _font_path()
            self._font_cache[pixels] = (
                ImageFont.truetype(path, pixels) if path else ImageFont.load_default()
            )
        return self._font_cache[pixels]

    def record_text(
        self,
        name: str,
        text: str,
        position: dict[str, float],
        role: str,
        font_pt: float,
        glyph_bbox: list[float],
    ) -> None:
        line_count = estimate_lines(text, position["width"], font_pt)
        resolved = round(font_pt * POINT_TO_LAYOUT_PIXEL, 4)
        required_height = line_count * resolved * 1.25 + 10
        payload = {
            "name": name,
            "bbox": [
                position["left"],
                position["top"],
                position["width"],
                position["height"],
            ],
            "glyph_bbox": glyph_bbox,
            "resolvedFontSize": resolved,
            "textLayout": {"lineCount": line_count},
            "overflow": role != "AUX" and required_height > position["height"] + 1e-6,
        }
        if (
            self.elements
            and self.elements[-1].get("name") == name
            and "textLayout" not in self.elements[-1]
        ):
            self.elements[-1].update(payload)
        else:
            self.elements.append(payload)

    def draw_text(
        self,
        text: str,
        position: dict[str, float],
        *,
        font_pt: float,
        color: str,
        alignment: str,
        vertical: str = "middle",
    ) -> list[float]:
        font = self.font(font_pt)
        left = int(round(position["left"] + 8))
        top = int(round(position["top"] + 4))
        width = max(1, int(round(position["width"] - 16)))
        height = max(1, int(round(position["height"] - 8)))
        wrapped = _wrap_for_preview(self.draw, text, font, width)
        bounds = self.draw.multiline_textbbox((0, 0), wrapped, font=font, spacing=2)
        text_width = bounds[2] - bounds[0]
        text_height = bounds[3] - bounds[1]
        if alignment == "center":
            x = left + max(0, (width - text_width) // 2)
        elif alignment == "right":
            x = left + max(0, width - text_width)
        else:
            x = left
        y = top + (max(0, (height - text_height) // 2) if vertical == "middle" else 0)
        self.draw.multiline_text(
            (x, y), wrapped, font=font, fill=f"#{color}", spacing=2, align=alignment
        )
        return [
            float(x + bounds[0]),
            float(y + bounds[1]),
            float(text_width),
            float(text_height),
        ]

    def add_box(
        self,
        name: str,
        position: dict[str, float],
        *,
        geometry: str = "roundRect",
        fill: str = COLORS["white"],
        line: str = COLORS["line"],
        line_width: float = 1.0,
        line_dash: str = "solid",
    ):
        from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        from pptx.util import Pt

        self.register_name(name)
        kind = {
            "rect": MSO_AUTO_SHAPE_TYPE.RECTANGLE,
            "roundRect": MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
            "ellipse": MSO_AUTO_SHAPE_TYPE.OVAL,
            "diamond": MSO_AUTO_SHAPE_TYPE.DIAMOND,
        }[geometry]
        shape = self.slide.shapes.add_shape(
            kind,
            px(position["left"]),
            px(position["top"]),
            px(position["width"]),
            px(position["height"]),
        )
        shape.name = name
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(line_width)
        if line_dash not in {"solid", "dash"}:
            raise ValueError(f"unknown line dash:{line_dash}")
        # Preserve the standard writer byte-for-byte when the optional POC
        # profile is absent.  python-pptx's implicit solid line is sufficient;
        # only the docket-ledger profile needs an explicit dash override.
        if line_dash == "dash":
            shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        bounds = (
            int(round(position["left"])),
            int(round(position["top"])),
            int(round(position["left"] + position["width"])),
            int(round(position["top"] + position["height"])),
        )
        if geometry == "ellipse":
            self.draw.ellipse(bounds, fill=f"#{fill}", outline=f"#{line}", width=max(1, int(line_width)))
        elif geometry == "diamond":
            left, top, right, bottom = bounds
            center_x, center_y = (left + right) // 2, (top + bottom) // 2
            self.draw.polygon(
                ((center_x, top), (right, center_y), (center_x, bottom), (left, center_y)),
                fill=f"#{fill}",
                outline=f"#{line}",
            )
        elif geometry == "roundRect":
            self.draw.rounded_rectangle(bounds, radius=12, fill=f"#{fill}", outline=f"#{line}", width=max(1, int(line_width)))
        else:
            self.draw.rectangle(bounds, fill=f"#{fill}", outline=f"#{line}", width=max(1, int(line_width)))
        self.elements.append(
            {
                "name": name,
                "bbox": [position["left"], position["top"], position["width"], position["height"]],
            }
        )
        return shape

    def add_rule(self, name: str, position: dict[str, float], color: str, width: float = 1.0):
        adjusted = dict(position)
        if adjusted["width"] <= 0:
            adjusted["width"] = max(1.0, width)
        if adjusted["height"] <= 0:
            adjusted["height"] = max(1.0, width)
        return self.add_box(name, adjusted, geometry="rect", fill=color, line=color, line_width=0.1)

    def add_text(
        self,
        name: str,
        text: str,
        position: dict[str, float],
        *,
        role: str,
        font_pt: float | None = None,
        bold: bool = False,
        color: str = COLORS["ink"],
        alignment: str = "left",
        word_wrap: bool = True,
    ):
        self.register_name(name)
        shape = self.slide.shapes.add_textbox(
            px(position["left"]), px(position["top"]), px(position["width"]), px(position["height"])
        )
        shape.name = name
        points = font_pt if font_pt is not None else TYPE_PT[role]
        _set_text_frame(
            shape.text_frame,
            str(text),
            font_pt=points,
            bold=bold,
            color=color,
            alignment=alignment,
            word_wrap=word_wrap,
        )
        glyph_bbox = self.draw_text(
            str(text),
            position,
            font_pt=points,
            color=color,
            alignment=alignment,
        )
        self.record_text(name, str(text), position, role, points, glyph_bbox)
        return shape


def _alignment(value: str):
    from pptx.enum.text import PP_ALIGN

    return {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[value]


def _ensure_run_fonts(run) -> None:
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.oxml.ns import qn

    properties = run._r.get_or_add_rPr()
    properties.set("lang", "zh-CN")
    font_nodes = []
    for tag in ("a:latin", "a:ea", "a:cs"):
        node = properties.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
        else:
            properties.remove(node)
        node.set("typeface", FONT)
        font_nodes.append(node)
    # CT_TextCharacterProperties is an ordered sequence: fill/effects precede
    # latin, which precedes ea, which precedes cs.  Appending this normalized
    # trio after python-pptx's fill node preserves that schema order.
    for node in font_nodes:
        properties.append(node)


def _set_text_frame(
    text_frame,
    text: str,
    *,
    font_pt: float,
    bold: bool,
    color: str,
    alignment: str,
    word_wrap: bool = True,
) -> None:
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
    from pptx.util import Pt

    text_frame.clear()
    text_frame.word_wrap = word_wrap
    text_frame.auto_size = MSO_AUTO_SIZE.NONE
    text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    text_frame.margin_left = px(8)
    text_frame.margin_right = px(8)
    text_frame.margin_top = px(4)
    text_frame.margin_bottom = px(4)
    for index, line in enumerate(str(text).split("\n")):
        paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
        paragraph.alignment = _alignment(alignment)
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(0)
        paragraph.line_spacing = 1.05
        run = paragraph.add_run()
        run.text = line
        run.font.name = FONT
        run.font.size = Pt(font_pt)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
        _ensure_run_fonts(run)


def _set_cell_border(cell, color: str = COLORS["line"], width: int = 12700) -> None:
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.oxml.ns import qn

    properties = cell._tc.get_or_add_tcPr()
    borders = []
    for edge in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        existing = properties.find(qn(edge))
        if existing is not None:
            properties.remove(existing)
        line = OxmlElement(edge)
        line.set("w", str(width))
        solid = OxmlElement("a:solidFill")
        shade = OxmlElement("a:srgbClr")
        shade.set("val", color)
        solid.append(shade)
        dash = OxmlElement("a:prstDash")
        dash.set("val", "solid")
        line.extend((solid, dash))
        borders.append(line)
    # CT_TableCellProperties requires lnL/lnR/lnT/lnB before the fill choice.
    # The fill is already present because callers style the cell first.
    for index, line in enumerate(borders):
        properties.insert(index, line)


def add_labeled_box(
    ctx: SlideContext,
    name: str,
    text: str,
    position: dict[str, float],
    *,
    role: str,
    geometry: str = "roundRect",
    fill: str = COLORS["white"],
    line: str = COLORS["line"],
    font_pt: float | None = None,
    bold: bool = False,
    color: str = COLORS["ink"],
    alignment: str = "left",
    word_wrap: bool = True,
    line_dash: str = "solid",
):
    shape = ctx.add_box(
        name,
        position,
        geometry=geometry,
        fill=fill,
        line=line,
        line_dash=line_dash,
    )
    # add_box has already reserved this final name; attach text without renaming.
    points = font_pt if font_pt is not None else TYPE_PT[role]
    _set_text_frame(
        shape.text_frame,
        str(text),
        font_pt=points,
        bold=bold,
        color=color,
        alignment=alignment,
        word_wrap=word_wrap,
    )
    glyph_bbox = ctx.draw_text(
        str(text),
        position,
        font_pt=points,
        color=color,
        alignment=alignment,
    )
    ctx.record_text(name, str(text), position, role, points, glyph_bbox)
    return shape


def semantic_id(lookup: dict[tuple[str, str, int], str], slide: dict[str, Any], slot: str, index: int) -> str:
    key = (slide["slide_id"], slot, index)
    if key not in lookup:
        raise ValueError(f"missing semantic inventory item {slide['slide_id']}/{slot}/{index}")
    return lookup[key]


def semantic_ids(lookup, slide: dict[str, Any], slot: str, items: list[Any]) -> list[str]:
    return [semantic_id(lookup, slide, slot, index) for index in range(len(items))]


def add_bullets(
    ctx: SlideContext,
    lookup,
    slide: dict[str, Any],
    prefix: str,
    slot: str,
    items: list[Any],
    frame: dict[str, float],
    *,
    font_pt: float = TYPE_PT["BODY"],
    row_height: float | None = None,
    marker_fill: str = COLORS["pale_blue"],
    marker_line: str = COLORS["blue"],
) -> None:
    height = row_height if row_height is not None else max(52, math.floor(frame["height"] / max(len(items), 1)))
    ids = semantic_ids(lookup, slide, slot, items)
    for index, item in enumerate(items):
        top = frame["top"] + index * height
        marker_position = {"left": frame["left"], "top": top + 8, "width": 26, "height": 26}
        add_labeled_box(
            ctx,
            text_role_name(f"{prefix}:marker:{index}", "AUX"),
            str(index + 1),
            marker_position,
            role="AUX",
            geometry="ellipse",
            fill=marker_fill,
            line=marker_line,
            bold=True,
            color=COLORS["blue"],
            alignment="center",
        )
        text = item if isinstance(item, str) else item["text"]
        ctx.add_text(
            text_role_name(semantic_name(f"{prefix}:item:{index}", [ids[index]]), "BODY"),
            text,
            {"left": frame["left"] + 40, "top": top, "width": frame["width"] - 40, "height": height - 4},
            role="BODY",
            font_pt=font_pt,
        )


def add_native_table(
    ctx: SlideContext,
    name: str,
    values: list[list[Any]],
    frame: dict[str, float],
    *,
    header: bool,
    semantic_values: list[str],
) -> None:
    from pptx.enum.text import MSO_ANCHOR

    if not values:
        return
    columns = max(len(row) for row in values)
    normalized = [[str(value or "") for value in row] + [""] * (columns - len(row)) for row in values]
    final_name = text_role_name(semantic_name(name, semantic_values), "BODY")
    ctx.register_name(final_name)
    shape = ctx.slide.shapes.add_table(
        len(normalized), columns, px(frame["left"]), px(frame["top"]), px(frame["width"]), px(frame["height"])
    )
    shape.name = final_name
    table = shape.table
    for column in table.columns:
        column.width = px(frame["width"] / columns)
    for row in table.rows:
        row.height = px(frame["height"] / len(normalized))
    for row_index, row in enumerate(normalized):
        for column_index, value in enumerate(row):
            cell = table.cell(row_index, column_index)
            cell.margin_left = px(6)
            cell.margin_right = px(6)
            cell.margin_top = px(3)
            cell.margin_bottom = px(3)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if header and row_index == 0:
                cell.fill.fore_color.rgb = rgb(COLORS["ink"])
                text_color = COLORS["white"]
                bold = True
            else:
                cell.fill.fore_color.rgb = rgb(COLORS["wash"] if row_index % 2 else "FAFBFD")
                text_color = COLORS["ink"]
                bold = False
            _set_cell_border(cell)
            _set_text_frame(
                cell.text_frame,
                value,
                font_pt=TYPE_PT["BODY"],
                bold=bold,
                color=text_color,
                alignment="left",
            )
    ctx.elements.append(
        {
            "name": final_name,
            "bbox": [frame["left"], frame["top"], frame["width"], frame["height"]],
            "resolvedFontSize": TYPE_PT["BODY"] * POINT_TO_LAYOUT_PIXEL,
            "table": {"rows": len(normalized), "columns": columns},
        }
    )
    left, top, width, height = (frame[key] for key in ("left", "top", "width", "height"))
    row_height = height / len(normalized)
    column_width = width / columns
    for row_index, row in enumerate(normalized):
        for column_index, value in enumerate(row):
            x0 = int(round(left + column_index * column_width))
            y0 = int(round(top + row_index * row_height))
            x1 = int(round(left + (column_index + 1) * column_width))
            y1 = int(round(top + (row_index + 1) * row_height))
            fill = COLORS["ink"] if header and row_index == 0 else (COLORS["wash"] if row_index % 2 else "FAFBFD")
            ctx.draw.rectangle((x0, y0, x1, y1), fill=f"#{fill}", outline=f"#{COLORS['line']}")
            ctx.draw_text(
                value,
                {"left": x0 + 2, "top": y0 + 2, "width": max(1, x1 - x0 - 4), "height": max(1, y1 - y0 - 4)},
                font_pt=TYPE_PT["BODY"],
                color=COLORS["white"] if header and row_index == 0 else COLORS["ink"],
                alignment="left",
            )


def add_chrome(ctx: SlideContext, lookup, slide: dict[str, Any], index: int, total: int, release_state: str) -> None:
    ctx.set_background(COLORS["white"])
    ctx.add_text(text_role_name(f"{slide['slide_id']}:eyebrow", "AUX"), DISPLAY_EYEBROW, {"left": PAGE["left"], "top": 30, "width": 430, "height": 24}, role="AUX", font_pt=10, bold=True, color=COLORS["blue"], word_wrap=False)
    ctx.add_text(text_role_name(semantic_name(f"{slide['slide_id']}:title", [semantic_id(lookup, slide, "title", 0)]), "SLIDE"), slide["title"], {"left": PAGE["left"], "top": 54, "width": 860, "height": 70}, role="SLIDE", bold=True)
    if release_state != "REVIEW_DRAFT":
        raise ValueError("客户可见状态仅允许审阅稿")
    add_labeled_box(ctx, text_role_name(f"{slide['slide_id']}:release-state", "MID"), DISPLAY_REVIEW_STATE, {"left": 956, "top": 36, "width": 260, "height": 90}, role="MID", fill=COLORS["pale_amber"], line="E8C890", bold=True, color=COLORS["amber"], alignment="center", word_wrap=False)
    ctx.add_rule(f"{slide['slide_id']}:header-rule", {"left": PAGE["left"], "top": 132, "width": PAGE["width"], "height": 0}, COLORS["line"])
    ctx.add_text(text_role_name(f"{slide['slide_id']}:page", "AUX"), f"{index + 1:02d} / {total:02d}", {"left": 1090, "top": 680, "width": 126, "height": 18}, role="AUX", font_pt=9, color=COLORS["muted"], alignment="right")


def render_cover(ctx: SlideContext, lookup, slide: dict[str, Any], profile: dict[str, Any]) -> None:
    ctx.set_background(COLORS["ink"])
    ctx.add_text(text_role_name(f"{slide['slide_id']}:series", "AUX"), DISPLAY_SERIES, {"left": 72, "top": 58, "width": 390, "height": 30}, role="AUX", font_pt=11, bold=True, color="9EC0FF", word_wrap=False)
    ctx.add_text(text_role_name(semantic_name(f"{slide['slide_id']}:matter", [semantic_id(lookup, slide, "title", 0), semantic_id(lookup, slide, "profile.matter_label", 0)]), "DECK"), profile["matter_label"], {"left": 72, "top": 138, "width": 1050, "height": 190}, role="DECK", bold=True, color=COLORS["white"])
    ctx.add_text(text_role_name(semantic_name(f"{slide['slide_id']}:client", [semantic_id(lookup, slide, "profile.client_label", 0)]), "MID"), profile["client_label"], {"left": 72, "top": 350, "width": 760, "height": 64}, role="MID", color="DCE7F8")
    ctx.add_text(text_role_name(semantic_name(f"{slide['slide_id']}:meta", [semantic_id(lookup, slide, "profile.materials_as_of", 0), semantic_id(lookup, slide, "profile.prepared_by", 0)]), "BODY"), f"材料截止：{profile['materials_as_of']}\n编制：{profile['prepared_by']}", {"left": 72, "top": 432, "width": 620, "height": 96}, role="BODY", color="B9C7D9")
    add_labeled_box(ctx, text_role_name(semantic_name(f"{slide['slide_id']}:draft", [semantic_id(lookup, slide, "profile.confidentiality", 0)]), "MID"), f"{DISPLAY_REVIEW_STATE}\n{display_confidentiality(profile['confidentiality'])}", {"left": 800, "top": 440, "width": 390, "height": 130}, role="MID", fill="24324A", line="52647F", bold=True, color="FFD999", alignment="center", word_wrap=False)
    ctx.add_text(text_role_name(f"{slide['slide_id']}:boundary", "AUX"), "零格式返工目标 · 不替代律师终审与客户放行", {"left": 72, "top": 650, "width": 900, "height": 26}, role="AUX", font_pt=10, color="8FA2BC")


def render_relationship(ctx, lookup, slide):
    parties = slide["content"]["parties"]
    chain = slide["content"]["chain"]
    if bool(parties) == bool(chain):
        raise ValueError(f"relationship页必须且只能包含parties或chain:{slide['slide_id']}")
    if len(parties) > 6 or len(chain) > 3:
        raise ValueError(f"relationship页容量超限:{slide['slide_id']}")
    if parties:
        columns = 2 if len(parties) == 4 else min(3, len(parties))
        rows = math.ceil(len(parties) / columns)
        gap_x, gap_y = 40, 30
        card_width = (1100 - gap_x * (columns - 1)) / columns
        card_height = 132 if rows == 1 else 150
        block_height = rows * card_height + max(0, rows - 1) * gap_y
        block_top = 148 + max(0, (490 - block_height) / 2)
        for index, party in enumerate(parties):
            row, column = divmod(index, columns)
            position = {"left": 90 + column * (card_width + gap_x), "top": block_top + row * (card_height + gap_y), "width": card_width, "height": card_height}
            add_labeled_box(ctx, text_role_name(semantic_name(f"{slide['slide_id']}:party:{index}", [semantic_id(lookup, slide, "parties", index)]), "BODY"), party["text"], position, role="BODY", geometry="rect", fill=COLORS["wash"], line="C5D1E2", bold=True, alignment="center")
        return
    area_top, area_height, row_gap = 148, 490, 18
    row_height = (area_height - row_gap * max(0, len(chain) - 1)) / len(chain)
    block_height = row_height * len(chain) + row_gap * max(0, len(chain) - 1)
    block_top = area_top + max(0, (area_height - block_height) / 2)
    for index, item in enumerate(chain):
        row_top = block_top + index * (row_height + row_gap)
        ctx.add_box(f"{slide['slide_id']}:chain-row:{index}", {"left": 78, "top": row_top, "width": 1112, "height": row_height}, geometry="rect", fill=COLORS["wash"] if index % 2 == 0 else "FAFBFD", line="CFD9E8")
        ctx.add_box(f"{slide['slide_id']}:chain-marker:{index}", {"left": 88, "top": row_top + row_height / 2 - 5, "width": 10, "height": 10}, geometry="ellipse", fill=COLORS["blue"], line=COLORS["blue"])
        ctx.add_text(text_role_name(semantic_name(f"{slide['slide_id']}:chain-text:{index}", [semantic_id(lookup, slide, "chain", index)]), "BODY"), item["text"], {"left": 112, "top": row_top + 12, "width": 1052, "height": row_height - 24}, role="BODY")


def _timeline_axis_positions(events: list[dict[str, Any]]) -> list[float]:
    """Map frozen ISO dates onto one honest horizontal scale without reordering."""

    if not events:
        return []
    parsed: list[date] = []
    for event in events:
        value = event.get("date")
        if not isinstance(value, str):
            raise ValueError("TIMELINE_DATE_SCALE_INVALID:date is not text")
        try:
            parsed_date = date.fromisoformat(value)
        except ValueError as exc:
            raise ValueError(f"TIMELINE_DATE_SCALE_INVALID:{value}") from exc
        if parsed_date.isoformat() != value:
            raise ValueError(f"TIMELINE_DATE_SCALE_NONCANONICAL:{value}")
        parsed.append(parsed_date)
    if parsed != sorted(parsed):
        raise ValueError("TIMELINE_DATE_SCALE_ORDER_DRIFT")
    if len(parsed) == 1:
        return [(TIMELINE_AXIS_LEFT + TIMELINE_AXIS_RIGHT) / 2]
    span = (parsed[-1] - parsed[0]).days
    if span <= 0:
        raise ValueError("TIMELINE_DATE_SCALE_DEGENERATE")
    width = TIMELINE_AXIS_RIGHT - TIMELINE_AXIS_LEFT
    return [
        TIMELINE_AXIS_LEFT + width * (value - parsed[0]).days / span
        for value in parsed
    ]


def _render_timeline_control(ctx, lookup, slide, events):
    """Retain the installed writer's rail/card composition as a same-IR control."""

    events = slide["content"]["events"]
    area_top = 140 if events else 154
    available_height = 638 - area_top
    row_gap = 16 if len(events) <= 3 else 10
    maximum = 180 if len(events) <= 2 else 140 if len(events) == 3 else 119
    fitted = (available_height - row_gap * max(0, len(events) - 1)) / len(events) if events else 0
    row_height = min(maximum, fitted)
    block_height = row_height * len(events) + row_gap * max(0, len(events) - 1) if events else 0
    block_top = area_top + max(0, (available_height - block_height) / 2)
    centers = [block_top + index * (row_height + row_gap) + row_height / 2 for index in range(len(events))]
    if len(centers) > 1:
        ctx.add_rule(f"{slide['slide_id']}:timeline-guide", {"left": TIMELINE_GUIDE_LEFT, "top": centers[0], "width": 0, "height": centers[-1] - centers[0]}, "A9BDE4", 2)
    for index, event in enumerate(events):
        row_top = block_top + index * (row_height + row_gap)
        # LibreOffice 26.8 renders the bold 16.5pt ISO date at about 107.27pt.
        # The former 105pt box had only 93pt after the two 6pt text margins, so
        # the last two digits wrapped.  This 129pt box leaves 117pt of real
        # text width.  The marker and body move together by only 16px, keeping
        # their original gaps and the 64px safe bound. ``word_wrap=False`` is
        # a second fail-safe; the OOXML verifier independently proves the box
        # is wide enough, so text cannot merely overflow into the marker.
        ctx.add_text(
            text_role_name(f"{slide['slide_id']}:event-date:{index}", "BODY"),
            event["date"],
            {
                "left": TIMELINE_DATE_LEFT,
                "top": row_top,
                "width": TIMELINE_DATE_WIDTH,
                "height": row_height,
            },
            role="BODY",
            bold=True,
            color=COLORS["blue"],
            alignment="right",
            word_wrap=False,
        )
        ctx.add_box(
            f"{slide['slide_id']}:event-marker:{index}",
            {"left": TIMELINE_MARKER_LEFT, "top": centers[index] - 11, "width": 22, "height": 22},
            geometry="ellipse",
            fill=COLORS["blue"],
            line=COLORS["blue"],
        )
        add_labeled_box(
            ctx,
            text_role_name(semantic_name(f"{slide['slide_id']}:event-text:{index}", [semantic_id(lookup, slide, "events", index)]), "BODY"),
            event["text"],
            {"left": TIMELINE_TEXT_LEFT, "top": row_top, "width": TIMELINE_TEXT_WIDTH, "height": row_height},
            role="BODY",
            geometry="rect",
            fill=COLORS["wash"] if index % 2 == 0 else "FAFBFD",
            line="CFD9E8",
        )


def _render_timeline_axis(ctx, lookup, slide, events):
    """Profile A: date-proportional horizontal axis with alternating labels."""

    positions = _timeline_axis_positions(events)
    if len(positions) > 1:
        ctx.add_rule(
            f"{slide['slide_id']}:timeline-guide",
            {
                "left": TIMELINE_AXIS_LEFT,
                "top": TIMELINE_AXIS_Y,
                "width": TIMELINE_AXIS_RIGHT - TIMELINE_AXIS_LEFT,
                "height": 0,
            },
            "7895BF",
            2,
        )
    for index, (event, marker_x) in enumerate(zip(events, positions)):
        above = index % 2 == 0
        label_left = marker_x - 150
        date_top = 194 if above else 446
        text_top = 240 if above else 492
        leader_top = 322 if above else TIMELINE_AXIS_Y
        leader_height = TIMELINE_AXIS_Y - 322 if above else 64
        ctx.add_text(
            text_role_name(f"{slide['slide_id']}:event-date:{index}", "BODY"),
            event["date"],
            {
                "left": label_left,
                "top": date_top,
                "width": 300,
                "height": 42,
            },
            role="BODY",
            bold=True,
            color=COLORS["blue"],
            alignment="center",
            word_wrap=False,
        )
        ctx.add_box(
            f"{slide['slide_id']}:event-marker:{index}",
            {
                "left": marker_x - 8,
                "top": TIMELINE_AXIS_Y - 8,
                "width": 16,
                "height": 16,
            },
            geometry="ellipse",
            fill=COLORS["white"],
            line=COLORS["blue"],
            line_width=2,
        )
        ctx.add_rule(
            f"{slide['slide_id']}:event-leader:{index}",
            {
                "left": marker_x,
                "top": leader_top,
                "width": 0,
                "height": leader_height,
            },
            "B5C3D8",
        )
        ctx.add_text(
            text_role_name(
                semantic_name(
                    f"{slide['slide_id']}:event-text:{index}",
                    [semantic_id(lookup, slide, "events", index)],
                ),
                "BODY",
            ),
            event["text"],
            {
                "left": label_left,
                "top": text_top,
                "width": 300,
                "height": 80,
            },
            role="BODY",
            font_pt=16.0,
            color=COLORS["ink"],
            alignment="center",
        )


def _render_timeline_register(ctx, lookup, slide, events):
    """Profile B: vertical docket register with a date index and thin rules."""

    fractions = [
        (value - TIMELINE_AXIS_LEFT)
        / (TIMELINE_AXIS_RIGHT - TIMELINE_AXIS_LEFT)
        for value in _timeline_axis_positions(events)
    ]
    row_centers = [231 + 300 * value for value in fractions]
    ctx.add_rule(
        f"{slide['slide_id']}:timeline-guide",
        {
            "left": 316,
            "top": 176,
            "width": 0,
            "height": 430,
        },
        "B8C2D0",
    )
    for index, (event, row_center) in enumerate(zip(events, row_centers)):
        content_top = row_center - 36
        divider_y = (
            (row_center + row_centers[index + 1]) / 2
            if index + 1 < len(row_centers)
            else min(606, row_center + 64)
        )
        ctx.add_text(
            text_role_name(f"{slide['slide_id']}:event-date:{index}", "BODY"),
            event["date"],
            {"left": 78, "top": content_top, "width": 210, "height": 72},
            role="BODY",
            bold=True,
            color=COLORS["blue"],
            alignment="right",
            word_wrap=False,
        )
        ctx.add_box(
            f"{slide['slide_id']}:event-marker:{index}",
            {"left": 311, "top": row_center - 5, "width": 10, "height": 10},
            geometry="rect",
            fill=COLORS["blue"],
            line=COLORS["blue"],
        )
        ctx.add_rule(
            f"{slide['slide_id']}:event-leader:{index}",
            {"left": 78, "top": divider_y, "width": 1112, "height": 0},
            "D7DEE8",
        )
        ctx.add_text(
            text_role_name(
                semantic_name(
                    f"{slide['slide_id']}:event-text:{index}",
                    [semantic_id(lookup, slide, "events", index)],
                ),
                "BODY",
            ),
            event["text"],
            {"left": 352, "top": content_top, "width": 820, "height": 72},
            role="BODY",
            font_pt=16.0,
            color=COLORS["ink"],
        )


def _render_undated_timeline(ctx, lookup, slide, milestones):
    if milestones:
        panel_top, panel_height = 154, 484
        ctx.add_box(f"{slide['slide_id']}:undated-panel", {"left": 78, "top": panel_top, "width": 1112, "height": panel_height}, fill=COLORS["pale_amber"], line="E8C890")
        ctx.add_text(text_role_name(f"{slide['slide_id']}:undated-title", "MID"), "未系日期事项", {"left": 98, "top": panel_top + 8, "width": 300, "height": 52}, role="MID", bold=True, color=COLORS["amber"])
        list_height = panel_height - 74
        add_bullets(ctx, lookup, slide, f"{slide['slide_id']}:undated", "undated_milestones", milestones, {"left": 112, "top": panel_top + 64, "width": 1048, "height": list_height}, row_height=math.floor(list_height / max(1, len(milestones))), marker_fill="FFE1A8", marker_line=COLORS["amber"])


def render_timeline(ctx, lookup, slide, timeline_profile: str | None = None):
    events = slide["content"]["events"]
    milestones = slide["content"]["undated_milestones"]
    if timeline_profile is None:
        _render_timeline_control(ctx, lookup, slide, events)
    elif timeline_profile == "a":
        _render_timeline_axis(ctx, lookup, slide, events)
    elif timeline_profile == "b":
        _render_timeline_register(ctx, lookup, slide, events)
    else:
        raise ValueError(f"unknown timeline profile:{timeline_profile}")
    _render_undated_timeline(ctx, lookup, slide, milestones)


def render_evidence(ctx, lookup, slide, table_headers):
    rows = slide["content"]["rows"]
    amounts = slide["content"]["amounts"]
    amount_notes = slide["content"]["amount_notes"]
    evidence = slide["content"]["evidence_index"]
    if sum(bool(value) for value in (rows, amounts, amount_notes, evidence)) != 1:
        raise ValueError(f"evidence页必须且只能包含rows/amounts/amount_notes/evidence_index之一:{slide['slide_id']}")
    frame = {"left": 72, "top": 154, "width": 1136, "height": 484}
    if rows:
        add_native_table(ctx, f"{slide['slide_id']}:matrix", [table_headers["matrix"], *[row["cells"] for row in rows]], frame, header=True, semantic_values=semantic_ids(lookup, slide, "rows", rows))
    elif amounts:
        add_native_table(ctx, f"{slide['slide_id']}:amounts", [table_headers["amounts"], *[row["cells"] for row in amounts]], frame, header=True, semantic_values=semantic_ids(lookup, slide, "amounts", amounts))
    elif amount_notes:
        add_bullets(ctx, lookup, slide, f"{slide['slide_id']}:amount-notes", "amount_notes", amount_notes, {"left": 90, "top": 166, "width": 1080, "height": 460}, row_height=math.floor(460 / len(amount_notes)))
    else:
        add_bullets(ctx, lookup, slide, f"{slide['slide_id']}:evidence", "evidence_index", evidence, {"left": 90, "top": 166, "width": 1080, "height": 460}, row_height=math.floor(460 / len(evidence)))


def render_stage_risk(ctx, lookup, slide, table_headers):
    stages, risks = slide["content"]["stages"], slide["content"]["risks"]
    if bool(stages) == bool(risks):
        raise ValueError(f"stage-risk页必须且只能包含stages或risks:{slide['slide_id']}")
    items, slot, header = (stages, "stages", table_headers["stages"]) if stages else (risks, "risks", table_headers["risks"])
    add_native_table(ctx, f"{slide['slide_id']}:{slot}", [header, *[row["cells"] for row in items]], {"left": 72, "top": 154, "width": 1136, "height": 484}, header=True, semantic_values=semantic_ids(lookup, slide, slot, items))


def render_next_steps(ctx, lookup, slide):
    items, gaps = slide["content"]["items"], slide["content"]["gaps"]
    if bool(items) == bool(gaps):
        raise ValueError(f"next-steps页必须且只能包含items或gaps:{slide['slide_id']}")
    if items:
        add_bullets(ctx, lookup, slide, f"{slide['slide_id']}:actions", "items", items, {"left": 90, "top": 166, "width": 1080, "height": 460}, row_height=math.floor(460 / len(items)))
    else:
        ctx.add_box(f"{slide['slide_id']}:gaps-panel", {"left": 78, "top": 154, "width": 1112, "height": 484}, fill=COLORS["wash"], line=COLORS["line"])
        add_bullets(ctx, lookup, slide, f"{slide['slide_id']}:gaps", "gaps", gaps, {"left": 104, "top": 176, "width": 1060, "height": 440}, row_height=math.floor(440 / len(gaps)), marker_fill=COLORS["pale_amber"], marker_line=COLORS["amber"])


def render_scope(ctx, lookup, slide, profile):
    items = slide["content"]["items"]
    add_bullets(ctx, lookup, slide, f"{slide['slide_id']}:scope", "items", items, {"left": 80, "top": 154, "width": 720, "height": 392}, row_height=72)
    add_labeled_box(ctx, text_role_name(f"{slide['slide_id']}:boundary-panel", "BODY"), f"材料截止\n{profile['materials_as_of']}\n\n状态\n{DISPLAY_REVIEW_STATE}\n\n边界\n零格式返工不替代事实、法律、策略、隐私与律师终审。", {"left": 842, "top": 154, "width": 350, "height": 392}, role="BODY", fill=COLORS["ink"], line=COLORS["ink"], color=COLORS["white"])
    ctx.add_text(text_role_name(f"{slide['slide_id']}:no-release", "BODY"), "未经指定律师终审与客户放行，不得标记为客户可交付或外发。", {"left": 80, "top": 574, "width": 1100, "height": 52}, role="BODY", bold=True, color=COLORS["amber"])


def render_slide(
    ctx,
    lookup,
    slide: dict[str, Any],
    index: int,
    total: int,
    ir: dict[str, Any],
    timeline_profile: str | None = None,
) -> None:
    table_headers = visible_table_headers(ir.get("source_schema"))
    if slide["role"] == "cover":
        render_cover(ctx, lookup, slide, ir["profile"])
    else:
        add_chrome(ctx, lookup, slide, index, total, ir["release_state"])
        role = slide["role"]
        if role == "executive_summary":
            add_bullets(ctx, lookup, slide, slide["slide_id"], "items", slide["content"]["items"], {"left": 90, "top": 160, "width": 1080, "height": 460}, row_height=88)
        elif role == "relationship":
            render_relationship(ctx, lookup, slide)
        elif role == "timeline":
            render_timeline(ctx, lookup, slide, timeline_profile)
        elif role == "evidence_matrix":
            render_evidence(ctx, lookup, slide, table_headers)
        elif role == "stage_risk":
            render_stage_risk(ctx, lookup, slide, table_headers)
        elif role == "next_steps":
            render_next_steps(ctx, lookup, slide)
        elif role == "scope":
            render_scope(ctx, lookup, slide, ir["profile"])
        else:
            raise ValueError(f"unsupported slide role {role}")
    notes = "\n".join([NOTES_SOURCE_HEADING, *slide["source_refs"], NOTES_BOUNDARY_HEADING, NOTES_BOUNDARY_TEXT])
    notes_frame = ctx.slide.notes_slide.notes_text_frame
    notes_frame.text = notes
    # notes_text_frame.text creates runs directly; route them through the same
    # zh-CN/typeface normalization the slide writers use so notes pages do not
    # fall back to the notes-master Latin theme font.
    for paragraph in notes_frame.paragraphs:
        for run in paragraph.runs:
            _ensure_run_fonts(run)


def write_qa(qa: Path, contexts: list[SlideContext]) -> None:
    from PIL import Image

    qa.mkdir(parents=True, exist_ok=True)
    for index, ctx in enumerate(contexts, 1):
        stem = f"slide-{index:02d}"
        ctx.image.save(qa / f"{stem}.png", format="PNG", optimize=False)
        (qa / f"{stem}.layout.json").write_text(
            json.dumps({"elements": ctx.elements}, ensure_ascii=False, sort_keys=True, indent=2) + "\n",
            encoding="utf-8",
        )
    columns = 4
    thumb_width, thumb_height = 320, 180
    rows = math.ceil(len(contexts) / columns)
    montage = Image.new("RGB", (columns * thumb_width, rows * thumb_height), "#E5E7EB")
    for index, ctx in enumerate(contexts):
        thumb = ctx.image.resize((thumb_width, thumb_height), Image.Resampling.LANCZOS)
        montage.paste(thumb, ((index % columns) * thumb_width, (index // columns) * thumb_height))
    montage.save(qa / "deck-montage.webp", format="WEBP", lossless=True, method=6)
    with (qa / "deck-inspect.ndjson").open("w", encoding="utf-8", newline="\n") as handle:
        for index, ctx in enumerate(contexts, 1):
            handle.write(json.dumps({"slide": index, "slide_id": ctx.slide_id, "elements": ctx.elements}, ensure_ascii=False, sort_keys=True) + "\n")


def write_presentation(
    input_path: Path,
    output_path: Path,
    qa: Path,
    *,
    timeline_profile: str | None = None,
) -> dict[str, Any]:
    dependency_report()
    from pptx import Presentation

    ir = json.loads(input_path.read_text(encoding="utf-8"))
    if timeline_profile is not None and timeline_profile not in TIMELINE_PROFILES:
        raise ValueError("timeline_profile must be a or b")
    if ir.get("schema") != "delivery-ir/1.0":
        raise ValueError("unsupported DeliveryIR")
    if ir.get("release_state") != "REVIEW_DRAFT":
        raise ValueError(
            "python-pptx writer is REVIEW_DRAFT-only; CLIENT_READY requires a separately reviewed human release-receipt path"
        )
    inventory = ir.get("semantic_inventory")
    if not isinstance(inventory, list) or not inventory:
        raise ValueError("missing semantic inventory")
    lookup: dict[tuple[str, str, int], str] = {}
    for item in inventory:
        key = (item["slide_id"], item["slot"], item["item_index"])
        if key in lookup:
            raise ValueError(f"duplicate semantic inventory locator {key}")
        lookup[key] = item["semantic_id"]

    presentation = Presentation()
    presentation.slide_width = px(WIDTH)
    presentation.slide_height = px(HEIGHT)
    # python-pptx updates cx/cy but retains ``screen4x3`` from its default
    # template.  Close all three attributes so Office never sees conflicting
    # 16:9 geometry and 4:3 presentation metadata.
    slide_size = presentation._element.sldSz
    slide_size.attrib.clear()
    slide_size.attrib.update(CLOSED_SLIDE_SIZE)
    _close_presentation_default_text_style(presentation)
    presentation.core_properties.author = "诉讼可视化交付工具"
    presentation.core_properties.last_modified_by = "诉讼可视化交付工具"
    presentation.core_properties.title = "诉讼可视化简报（审阅稿）"
    presentation.core_properties.subject = "原生可编辑客户审阅材料"
    presentation.core_properties.comments = "审阅稿；须经人工法律复核与客户放行。"
    contexts: list[SlideContext] = []
    blank = presentation.slide_layouts[6]
    for index, ir_slide in enumerate(ir["slides"]):
        slide = presentation.slides.add_slide(blank)
        ctx = SlideContext(slide, ir_slide["slide_id"])
        render_slide(
            ctx,
            lookup,
            ir_slide,
            index,
            len(ir["slides"]),
            ir,
            timeline_profile,
        )
        contexts.append(ctx)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    fd, temp_name = tempfile.mkstemp(prefix=f".{output_path.name}.", suffix=".tmp", dir=output_path.parent)
    os.close(fd)
    temp = Path(temp_name)
    try:
        presentation.save(temp)
        with temp.open("rb") as handle:
            os.fsync(handle.fileno())
        os.replace(temp, output_path)
        write_qa(qa, contexts)
    finally:
        if temp.exists():
            temp.unlink()
    return {"ok": True, "engine": "python-pptx", "engine_version": EXPECTED_PYTHON_PPTX, "slide_count": len(ir["slides"]), "output": output_path.name, "timeline_profile": timeline_profile}


def parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--doctor", action="store_true")
    parser.add_argument("--input")
    parser.add_argument("--output")
    parser.add_argument("--qa")
    parser.add_argument("--timeline-profile", choices=("a", "b"))
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv or sys.argv[1:])
    try:
        report = dependency_report()
        if args.doctor:
            print(json.dumps(report, ensure_ascii=False, sort_keys=True))
            return 0
        if not args.input or not args.output or not args.qa:
            raise ValueError("--input/--output/--qa are required")
        result = write_presentation(
            Path(args.input),
            Path(args.output),
            Path(args.qa),
            timeline_profile=args.timeline_profile,
        )
        print(json.dumps(result, ensure_ascii=False, sort_keys=True))
        return 0
    except DependencyUnavailable as exc:
        print(f"HOLD-PPTX-DEPENDENCY:{exc}", file=sys.stderr)
        return EXIT_HOLD
    except (OSError, ValueError, KeyError, TypeError) as exc:
        print(f"PPTX-WRITE-FAILED:{type(exc).__name__}:{exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
